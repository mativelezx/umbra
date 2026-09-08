"""Read-only cross-check of the thesis, source artifacts and rendered PDF.

This checks explicit contracts, not academic approval or psychological validity.
"""
import argparse
import csv
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pypdf import PdfReader

parser = argparse.ArgumentParser()
parser.add_argument('docx', type=Path)
parser.add_argument('pdf', type=Path)
parser.add_argument('--repo', type=Path, required=True)
parser.add_argument('--manifest', type=Path, required=True)
parser.add_argument('--pages', type=Path, required=True)
parser.add_argument('--output', type=Path, required=True)
args = parser.parse_args()
doc = Document(args.docx)
pdf = PdfReader(args.pdf)
manifest = json.loads(args.manifest.read_text())
pages = json.loads(args.pages.read_text())
checks = []


def check(group, name, condition, detail=''):
    checks.append({'control': group, 'check': name, 'passed': bool(condition), 'detail': detail})


def read(relative):
    return (args.repo / relative).read_text()


def es_number(number, digits=3):
    return f'{number:.{digits}f}'.replace('.', ',')


def reported_metric_matches(cell, value):
    normalized = cell.strip().replace('−', '-')
    if normalized == '<0,001':
        return 0 < value < 0.001
    return normalized == es_number(value)


text = '\n'.join(page.extract_text() or '' for page in pdf.pages).replace('\u200b', '')
flat = re.sub(r'\s+', ' ', text)
outline = [title for title, _ in manifest['outline']]
required = ['Resumen', 'Abstract', 'Introducción', 'Justificación',
            'Objetivo general del proyecto', 'Objetivos específicos del proyecto',
            'Marco teórico referencial', 'Diseño metodológico', 'Relevamiento',
            'Procesos de negocio', 'Diagnóstico y propuesta',
            'Objetivo, límites y alcances del prototipo', 'Descripción del sistema',
            'Seguridad', 'Análisis de costos', 'Análisis de riesgos',
            'Conclusiones', 'Demo', 'Referencias', 'Anexos']
for title in required:
    check('1-consignas', title, title in outline and title in pages)
check('1-consignas', 'Demo entre conclusiones y referencias',
      outline.index('Conclusiones') < outline.index('Demo') < outline.index('Referencias'))
check('1-consignas', 'Resumen y abstract entre 150 y 250 palabras',
      all(150 <= n <= 250 for n in manifest['summary_words']), str(manifest['summary_words']))
check('1-consignas', 'Resumen y abstract en páginas separadas', pages['Resumen'] != pages['Abstract'])
check('1-consignas', 'Tres capturas del proceso núcleo', len(manifest['captions']['Figura']) >= 8)
check('1-consignas', 'Scrum: backlog, historias y primer sprint',
      all(item in outline for item in ['Product Backlog', 'Historias de Usuario', 'Sprint Backlog', 'Detalle del primer sprint']))
check('1-consignas', 'DER, diccionario y arquitectura',
      all(item in outline for item in ['Estructura de datos', 'Diccionario de datos', 'Diagrama de arquitectura']))
check('1-consignas', 'Fuentes y limitación del cronograma',
      all(item in flat for item in ['Cronograma reconstruido', 'Git', 'Canvas', 'no equivalen a horas trabajadas']))
check('1-consignas', 'Diferencia entre fechas históricas explicitada',
      'La secuencia exacta de esa decisión no quedó acreditada' in flat)

package = json.loads(read('package.json'))
check('2-codigo-modelo', 'Versión Next.js coincide', 'Next.js ' + package['dependencies']['next'] in flat)
check('2-codigo-modelo', 'Runtimes largos de seis rutas de IA', all(
    "runtime = 'nodejs'" in read(f'app/api/{route}/route.ts')
    and 'maxDuration = 120' in read(f'app/api/{route}/route.ts')
    for route in ['analyze', 'plan', 'narrative', 'chat', 'onboarding/next', 'onboarding/seed']))
metrics = json.loads(read('ml/eval_metrics.json'))
comparison = json.loads(read('ml/comparison-2026-09-07.json'))
dims = ['openness', 'conscientiousness', 'extraversion', 'agreeableness', 'neuroticism']
check('2-codigo-modelo', 'Cinco dimensiones insuficientes',
      all(metrics['blocks']['combined']['per_dimension_status'][d] == 'low_confidence' for d in dims)
      and 'cinco dimensiones' in flat and 'evidencia insuficiente' in flat)
rmse = doc.tables[28]
for row_number, dim in enumerate(dims, 1):
    for col, model in enumerate(['constant_training_mean', 'shipped_tfidf_ridge', 'shipped_distilbert_ridge'], 1):
        expected = es_number(comparison['results'][dim][model]['rmse_on_binary_labels'], 4)
        check('2-codigo-modelo', f'Tabla 29 {dim} {model}', rmse.cell(row_number, col).text == expected)
    for block, table_index in [('english_only', 29), ('combined', 30)]:
        record = metrics['blocks'][block]['metrics'][dim]
        table = doc.tables[table_index]
        for col, key in [(2, 'r2'), (3, 'r')]:
            check('2-codigo-modelo', f'Métricas {block} {dim} {key}',
                  reported_metric_matches(table.cell(row_number, col).text, record[key]))
        mse_expected = f"{record['mse']:,.2f}".replace(',', '_').replace('.', ',').replace('_', '.')
        check('2-codigo-modelo', f'MSE {block} {dim}', table.cell(row_number, 1).text == mse_expected)
splits = {}
for name, count in [('train', 1989), ('val', 248), ('test', 250)]:
    with (args.repo / f'ml/data/splits/{name}.csv').open() as file:
        rows = list(csv.DictReader(file))
    splits[name] = {row['id'] for row in rows}
    check('2-codigo-modelo', 'Partición ' + name, len(rows) == count and len(splits[name]) == count)
check('2-codigo-modelo', 'IDs de particiones sin solapamiento', not (
    splits['train'] & splits['test'] or splits['train'] & splits['val'] or splits['val'] & splits['test']))
for relative, expected in comparison['sha256'].items():
    check('2-codigo-modelo', 'Artefacto sin cambios: ' + relative,
          hashlib.sha256((args.repo / 'ml' / relative).read_bytes()).hexdigest() == expected)
bfi = read('lib/assessment/bfi2s.ts')
check('2-codigo-modelo', 'BFI-2-S: treinta ítems, 1–5 e inversión', all(
    item in bfi for item in ['.length(30)', '.min(1).max(5)', '6 - value', ' / 6']))
route = read('app/api/self-report/route.ts')
check('2-codigo-modelo', 'Autoinforme separado del entrenamiento',
      'createSelfReport(input.answers' in route and '...previous, selfReport' in route
      and '.eq(\'user_id\', user.id)' in route and 'researchUse: false' in bfi)
sql = '\n'.join(file.read_text() for file in (args.repo / 'supabase/migrations').glob('*.sql'))
tables = set(re.findall(r'CREATE TABLE (?:IF NOT EXISTS )?public\.(\w+)', sql, re.I))
check('2-codigo-modelo', 'Quince tablas públicas en migraciones', len(tables) == 15, ', '.join(sorted(tables)))
for table in sorted(tables):
    check('2-codigo-modelo', 'RLS declarado: ' + table,
          bool(re.search(r'ALTER TABLE public\.' + table + r' ENABLE ROW LEVEL SECURITY', sql, re.I)))
retention = read('supabase/migrations/20260907020000_preserve_profile_results_on_purge.sql')
check('2-codigo-modelo', 'Retención conserva autoinforme y ML por updated_at', all(
    item in retention for item in ["'selfReport'", "'ml'", "updated_at < NOW() - INTERVAL '30 days'"]))
check('2-codigo-modelo', 'Límites remotos declarados', all(
    item in flat for item in ['dominio propio verificado', 'no el transcurso real de treinta días', 'corte global de consumo']))

check('3-archivo-final', 'Noventa páginas sin portada numerada', len(pdf.pages) == 90)
check('3-archivo-final', 'Todos los destinos del índice localizados',
      set(outline).union(*manifest['captions'].values()) == set(pages))
for n, section in enumerate(doc.sections):
    check('3-archivo-final', f'Márgenes de 3 cm, sección {n+1}', all(
        abs(value.cm - 3) < 0.01 for value in [section.top_margin, section.bottom_margin, section.left_margin, section.right_margin]))
    check('3-archivo-final', f'A4, sección {n+1}', sorted([round(section.page_width.cm, 1), round(section.page_height.cm, 1)]) == [21.0, 29.7])
for name, size, bold, italic in [('Normal', 12, False, False), ('Heading 1', 14, True, False), ('Heading 2', 12, False, True), ('Heading 3', 12, False, True)]:
    style = doc.styles[name]
    check('3-archivo-final', 'Tipografía: ' + name, style.font.name == 'Times New Roman' and style.font.size.pt == size and style.font.bold == bold and style.font.italic == italic)
check('3-archivo-final', 'Interlineado y sangría del cuerpo',
      doc.styles['Normal'].paragraph_format.line_spacing == 1.5 and abs(doc.styles['Normal'].paragraph_format.first_line_indent.cm - 1.27) < 0.01)
check('3-archivo-final', 'Encabezados arriba a la derecha según rúbrica CAE',
      all(s.header.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT for s in doc.sections))
check('3-archivo-final', 'Tablas con encabezado repetible', all(t.rows[0]._tr.find('.//' + qn('w:tblHeader')) is not None for t in doc.tables))
links = [rel.target_ref for rel in doc.part.rels.values() if rel.reltype.endswith('/hyperlink')]
check('3-archivo-final', 'Enlace de demo exacto', 'https://umbra-sigma.vercel.app' in links)
check('3-archivo-final', 'Enlace de rama exacto', 'https://github.com/mativelezx/umbra/tree/codex/reentrega-final-2026-09-07' in links)
check('3-archivo-final', 'Enlaces sin punto de oración agregado', not any(link.endswith(('.', ',', ';')) for link in links))
check('3-archivo-final', 'Sin marcadores de borrador o rutas privadas', not any(
    marker in flat for marker in ['COMPLETAR', '[PENDIENTE', '/private/tmp/', '/Users/', '15.5.18', 'sk-ant-']))
report = {'checks': checks, 'passed': sum(c['passed'] for c in checks),
          'failed': [c for c in checks if not c['passed']],
          'docx_sha256': hashlib.sha256(args.docx.read_bytes()).hexdigest(),
          'pdf_sha256': hashlib.sha256(args.pdf.read_bytes()).hexdigest(),
          'limits': ['No certifica aprobación del CAE ni autoría o comprensión del estudiante.',
                     'No valida psicológicamente el ML ni sustituye una evaluación con participantes.',
                     'Correo remoto, restauración de respaldos y corte global pendientes.',
                     'La cronología reconstruida requiere confirmación del autor.']}
args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps({key: report[key] for key in ['passed', 'failed', 'docx_sha256', 'pdf_sha256']}, ensure_ascii=False))
raise SystemExit(1 if report['failed'] else 0)
